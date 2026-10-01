"""Write a copy of the user's draft with problem citations highlighted and explained in Word comments.

Called by build_report.py (after every quote has been re-verified), or standalone:
    python annotate_draft.py results.verified.json draft.docx -o draft_批注版.docx

For each item whose status is not "supported":
  - put 【原句 Original】 in front of the sentence and, if there is a rewrite, 【改写句 Rewrite】 + the
    rewrite (green italic) right after it
  - highlight `problem_phrase` (or the whole sentence if no phrase is given) in the status colour
  - attach a bilingual comment: status, issue (中文 + English), verified source quote + page, suggested rewrite
  - if `sentence_zh` is given (bilingual drafts), do the same for the Chinese sentence, using
    `problem_phrase_zh` and `suggested_rewrite_zh`
The original draft file is never modified.
"""
import argparse
import copy
import json
import re
import sys
from pathlib import Path

from docx.enum.text import WD_COLOR_INDEX
from docx.oxml.ns import qn
from docx.shared import RGBColor
from docx.text.run import Run

sys.stdout.reconfigure(encoding="utf-8")

COLOR = {
    "contradicted": WD_COLOR_INDEX.RED,
    "partial": WD_COLOR_INDEX.YELLOW,
    "not_found": WD_COLOR_INDEX.GRAY_25,
    "source_missing": WD_COLOR_INDEX.GRAY_25,
}
LABEL = {
    "contradicted": "🔴 与原文不一致 / Contradicts the source",
    "partial": "🟡 原文只部分支持 / Only partially supported",
    "not_found": "⚪ 未定位到原文，请自己确认 / Not located in the source — please check",
    "source_missing": "⚫ 找不到这篇文献的 PDF / Source PDF not available",
}


def _norm_map(text):
    """Whitespace/quote-insensitive version of text, with index map back to the original."""
    out, idx = [], []
    for i, c in enumerate(text):
        c = {"’": "'", "‘": "'", "“": '"', "”": '"', "–": "-", "—": "-", " ": " "}.get(c, c)
        if c.isspace():
            if out and out[-1] == " ":
                continue
            c = " "
        out.append(c)
        idx.append(i)
    return "".join(out), idx


def find_span(text, needle):
    if not needle:
        return None
    pos = text.find(needle)
    if pos >= 0:
        return pos, pos + len(needle)
    nt, it = _norm_map(text)
    nn, _ = _norm_map(needle.strip())
    pos = nt.find(nn)
    if pos < 0:
        pos = nt.lower().find(nn.lower())
    if pos < 0:
        return None
    return it[pos], it[pos + len(nn) - 1] + 1


def split_runs(paragraph, start, end):
    """Split runs so that [start, end) of the paragraph's run text is covered by whole runs; return them."""
    selected, offset = [], 0
    for run in list(paragraph.runs):
        t = run.text
        a, b = offset, offset + len(t)
        offset = b
        if b <= start or a >= end or not t:
            continue
        s, e = max(start, a) - a, min(end, b) - a
        if s > 0:
            pre = copy.deepcopy(run._r)
            run._r.addprevious(pre)
            type(run)(pre, paragraph).text = t[:s]
        if e < len(t):
            post = copy.deepcopy(run._r)
            run._r.addnext(post)
            type(run)(post, paragraph).text = t[e:]
        run.text = t[s:e]
        selected.append(run)
    return selected


def locate(doc, needle):
    for p in doc.paragraphs:
        text = "".join(r.text for r in p.runs)
        span = find_span(text, needle)
        if span:
            return p, span
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    text = "".join(r.text for r in p.runs)
                    span = find_span(text, needle)
                    if span:
                        return p, span
    return None, None


LABEL_ORIGINAL = "【原句 Original】"
LABEL_REWRITE = "【改写句 Rewrite】"
REWRITE_COLOR = RGBColor(0x00, 0x80, 0x3C)
NO_REWRITE = "(none — the source does not support this; delete the sentence or find another source / 无：原文不支持，建议删除或另找来源)"
NO_REWRITE_ZH = "（无：原文不支持，建议删除或另找来源）"


def _new_run(template, paragraph, text, *, before=None, after=None):
    """Insert a plain run next to `template`, inheriting its font but not highlight/comment formatting."""
    r = copy.deepcopy(template._r)
    for child in list(r):
        if child.tag != qn("w:rPr"):
            r.remove(child)
    (template._r.addprevious if before else template._r.addnext)(r)
    run = Run(r, paragraph)
    run.text = text
    run.font.highlight_color = None
    return run


def mark(doc, sentence, phrase, color, comment_lines, rewrite=""):
    """Label the sentence as 原句, highlight the problem phrase (or whole sentence), append the 改写句,
    and anchor the comment on the highlighted text."""
    p, span = locate(doc, sentence)
    if not p and phrase:  # sentence not copied exactly — fall back to marking just the phrase
        p, span = locate(doc, phrase)
        rewrite, labelled = "", False
    else:
        labelled = True
    if not p:
        return False
    s0, s1 = span
    target = (s0, s1)
    if phrase and labelled:
        text = "".join(r.text for r in p.runs)[s0:s1]
        ps = find_span(text, phrase)
        if ps:
            target = (s0 + ps[0], s0 + ps[1])
    hl = split_runs(p, *target)
    if not hl:
        return False
    for r in hl:
        r.font.highlight_color = color
    if labelled:
        sent_runs = split_runs(p, s0, s1)
        lab = _new_run(sent_runs[0], p, LABEL_ORIGINAL, before=True)
        lab.bold = True
        if rewrite:
            last = sent_runs[-1]
            gap = "" if re.search(r"[一-鿿]", sentence) else " "
            rw = _new_run(last, p, gap + rewrite, after=True)
            rw.italic, rw.bold = True, False
            rw.font.color.rgb = REWRITE_COLOR
            rl = _new_run(last, p, gap + LABEL_REWRITE, after=True)
            rl.bold = True
            rl.font.color.rgb = REWRITE_COLOR
    c = doc.add_comment(hl, text=comment_lines[0], author="Citation Check", initials="CC")
    for line in comment_lines[1:]:
        c.add_paragraph(line)
    return True


def comment_for(item):
    lines = [f"[#{item['id']}] {LABEL[item['overall']]}", f"引用 Citation: {item.get('citation', '')}"]
    for n, c in enumerate(item.get("claims", []), 1):
        prefix = f"({n}) " if len(item["claims"]) > 1 else ""
        if c["status"] == "supported":
            lines.append(f"{prefix}✓ 有原文支持 / Supported: {c['claim']}")
        elif c.get("issue_zh"):
            lines.append(f"{prefix}问题：{c['issue_zh']}")
        if c.get("issue_en") and c["status"] != "supported":
            lines.append(f"{prefix}Issue: {c['issue_en']}")
        for ev in c.get("evidence", []):
            lines.append(f"{prefix}原文 Source (PDF p.{'-'.join(map(str, ev['pages']))}): “{ev['pdf_text']}”")
    if item.get("suggested_rewrite"):
        lines.append(f"建议改写 Suggested rewrite: {item['suggested_rewrite']}")
    if item.get("suggested_rewrite_zh"):
        lines.append(f"中文改写：{item['suggested_rewrite_zh']}")
    return lines


def annotate(results, draft, out):
    import docx
    doc = docx.Document(str(draft))
    done, missed = 0, []
    for item in sorted(results["items"], key=lambda i: i["id"]):
        status = item.get("overall")
        if status not in COLOR:
            continue
        lines = comment_for(item)
        # partial/contradicted with no rewrite = the source cannot support any version of this sentence
        no_rewrite = NO_REWRITE if status in ("partial", "contradicted") else ""
        if mark(doc, item["sentence"], item.get("problem_phrase"), COLOR[status], lines,
                item.get("suggested_rewrite") or no_rewrite):
            done += 1
        else:
            missed.append(item["id"])
        if item.get("sentence_zh"):
            mark(doc, item["sentence_zh"], item.get("problem_phrase_zh"), COLOR[status],
                 [f"[#{item['id']}] 对应英文句的问题见英文批注 / See the comment on the English sentence #{item['id']}"],
                 item.get("suggested_rewrite_zh") or (NO_REWRITE_ZH if no_rewrite else ""))
    doc.save(str(out))
    return done, missed


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("results")
    ap.add_argument("draft")
    ap.add_argument("-o", "--out")
    args = ap.parse_args()
    results = json.loads(Path(args.results).read_text(encoding="utf-8"))
    out = args.out or str(Path(args.draft).with_name(Path(args.draft).stem + "_批注版.docx"))
    done, missed = annotate(results, args.draft, out)
    print(f"annotated draft -> {out} | marked {done} | could not locate: {missed}")


if __name__ == "__main__":
    main()

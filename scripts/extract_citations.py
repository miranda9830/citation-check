"""Extract English sentences that carry APA-style in-text citations from a draft.

Usage:
    python extract_citations.py draft.docx -o claims.json
    python extract_citations.py draft.txt  -o claims.json

Handles bilingual (Chinese/English) drafts: sentences that are mostly CJK are skipped,
because the English text is what gets submitted and what must match the source.
Also parses the reference list (after a "References"/"参考文献" heading, or any
paragraph that looks like an APA reference entry) so titles can disambiguate sources.
"""
import argparse
import json
import re
import sys
from pathlib import Path

sys.stdout.reconfigure(encoding="utf-8")

NAME = r"[A-Z][A-Za-z'’\-]+(?:\s(?:van|von|de|der|den|la|le|du|di)\s[A-Z][A-Za-z'’\-]+)?"
YEAR = r"(?:19|20)\d{2}[a-z]?|n\.d\.|in press"
# Narrative: "Vial (2019)", "Verhoef et al. (2021)", "Gong and Ribiere (2021, p. 3)"
NARRATIVE = re.compile(
    rf"({NAME}(?:\s(?:and|&)\s{NAME})?(?:\set\sal\.?)?)(?:'s|’s)?\s\(({YEAR})(?:,\s*(pp?\.\s*[\d\-–]+))?\)"
)
PAREN = re.compile(r"\(([^()]*?(?:(?:19|20)\d{2}[a-z]?|n\.d\.)[^()]*?)\)")
# One citation inside a parenthetical: "Vial, 2019" / "Hanelt et al., 2021, p. 5" / "Nambisan et al., 2017, 2019"
PAREN_ITEM = re.compile(
    rf"(?:e\.g\.,?\s*|see\s+(?:also\s+)?|cf\.\s*)?({NAME}(?:(?:,\s{NAME})*,?\s(?:and|&)\s{NAME})?(?:\set\sal\.?)?),\s*((?:(?:{YEAR})(?:,\s*)?)+)(?:,\s*(pp?\.\s*[\d\-–]+))?"
)
REF_ENTRY = re.compile(rf"^({NAME}),\s+(?:[A-Z]\.\s?)+.*?\(((?:19|20)\d{{2}}[a-z]?|n\.d\.)[^)]*\)\.\s*(.+?)[.?!]\s")
ABBREV = ["et al.", "e.g.", "i.e.", "cf.", "vs.", "p.", "pp.", "Vol.", "No.", "Fig.", "approx."]


def cjk_ratio(s):
    letters = [c for c in s if c.isalpha()]
    if not letters:
        return 0
    return sum(1 for c in letters if "一" <= c <= "鿿") / len(letters)


def split_sentences(text):
    protected = text
    for i, a in enumerate(ABBREV):
        protected = protected.replace(a, a.replace(".", f"§{i}§"))
    parts = re.split(r"(?<=[.!?])\s+(?=[A-Z“\"(])|(?<=[。！？])", protected)
    out = []
    for p in parts:
        for i, a in enumerate(ABBREV):
            p = p.replace(a.replace(".", f"§{i}§"), a)
        if p.strip():
            out.append(p.strip())
    return out


def read_paragraphs(path):
    p = Path(path)
    if p.suffix.lower() == ".docx":
        import docx
        d = docx.Document(str(p))
        paras = [para.text for para in d.paragraphs]
        for t in d.tables:
            for row in t.rows:
                for cell in row.cells:
                    paras.append(cell.text)
        return paras
    return p.read_text(encoding="utf-8", errors="replace").splitlines()


def parse_citations(sentence):
    cites = []
    for m in NARRATIVE.finditer(sentence):
        cites.append({"authors": m.group(1), "year": m.group(2), "pages": m.group(3), "raw": m.group(0), "form": "narrative"})
    for m in PAREN.finditer(sentence):
        inner = m.group(1)
        # skip the year-only parenthetical of a narrative citation, already captured
        if re.fullmatch(rf"\s*(?:{YEAR})(?:,\s*pp?\.\s*[\d\-–]+)?\s*", inner):
            continue
        for chunk in inner.split(";"):
            im = PAREN_ITEM.search(chunk.strip())
            if not im:
                continue
            years = [y for y in re.split(r",\s*", im.group(2)) if re.fullmatch(YEAR, y.strip())]
            for y in years:
                cites.append({"authors": im.group(1), "year": y.strip(), "pages": im.group(3), "raw": chunk.strip(), "form": "parenthetical"})
    for c in cites:
        c["first_author"] = re.split(r"\s(?:and|&|et)\s|,", c["authors"])[0].strip()
        c["key"] = f"{c['first_author']}|{re.sub(r'[a-z]$', '', c['year'])}"
    return cites


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("draft")
    ap.add_argument("-o", "--out", default="claims.json")
    args = ap.parse_args()

    paras = read_paragraphs(args.draft)
    items, refs = [], []
    in_refs = False
    for pi, para in enumerate(paras):
        text = para.strip()
        if not text:
            continue
        if re.fullmatch(r"(References?|Reference List|Bibliography|参考文献)\s*[:：]?", text, re.I):
            in_refs = True
            continue
        rm = REF_ENTRY.match(text + " ")
        if in_refs or rm:
            if rm:
                refs.append({"raw": text, "first_author": rm.group(1), "year": rm.group(2), "title": rm.group(3).strip()})
            continue
        for s in split_sentences(text):
            if cjk_ratio(s) > 0.3:
                continue
            cites = parse_citations(s)
            if cites:
                items.append({"id": len(items) + 1, "para": pi, "sentence": s, "citations": cites})

    out = {"draft": str(Path(args.draft).resolve()), "items": items, "references": refs,
           "unique_sources": sorted({c["key"] for it in items for c in it["citations"]})}
    Path(args.out).write_text(json.dumps(out, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"{len(items)} citing sentences, {len(out['unique_sources'])} unique sources, {len(refs)} reference entries -> {args.out}")
    for k in out["unique_sources"]:
        print("  ", k)


if __name__ == "__main__":
    main()
